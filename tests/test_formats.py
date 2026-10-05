"""Multi-format support: intake, protected release, and leak tracing for every accepted file type."""
import io, pytest
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from tracevault import crypto as C, formats as F
from tracevault.api import create_app

USERS = {"adm": {"password": "adminpass1", "role": "admin"}, "bob": {"password": "bobpass123", "role": "submitter"}}
needs_office = pytest.mark.skipif(not F.soffice(), reason="LibreOffice not installed")

@pytest.fixture(scope="module")
def c(tmp_path_factory):
    cl = TestClient(create_app(str(tmp_path_factory.mktemp("f")), USERS))
    adm = H(cl, "adm")
    assert cl.post("/recipient/enroll", headers=H(cl, "bob"), json={"display_name": "Bob Rivera", "passphrase": "bob-keystore-pass"}).status_code == 200
    return cl

def H(c, u):
    return {"Authorization": "Bearer " + c.post("/auth/login", json={"username": u, "password": USERS[u]["password"]}).json()["token"]}

def png_bytes(fmt="PNG", size=(900, 700)):
    im = Image.new("RGB", size, (235, 238, 244)); d = ImageDraw.Draw(im)
    for i in range(0, size[1], 40):
        d.text((30, i + 10), f"Confidential line {i} lorem ipsum dolor sit amet", fill=(30, 30, 60))
        d.rectangle([20, i, size[0] - 20, i + 2], fill=(150, 160, 190))
    b = io.BytesIO(); im.save(b, fmt); return b.getvalue()

def docx_bytes():
    import docx
    d = docx.Document(); d.add_heading("Operation plan", 1)
    for i in range(30): d.add_paragraph(f"Paragraph {i}: confidential material for restricted distribution only.")
    b = io.BytesIO(); d.save(b); return b.getvalue()

def pptx_bytes():
    import pptx
    p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = "Secret roadmap"
    s.placeholders[1].text = "Q4 launch plans"
    b = io.BytesIO(); p.save(b); return b.getvalue()

def xlsx_bytes():
    import openpyxl
    w = openpyxl.Workbook(); ws = w.active
    for r in range(1, 40): ws.append([f"Item {r}", r * 11, "restricted"])
    b = io.BytesIO(); w.save(b); return b.getvalue()

def test_detection_by_content_not_name():
    assert F.detect(png_bytes(), "x.txt") == "png"
    assert F.detect(png_bytes("JPEG"), "x.png") == "jpeg"
    assert F.detect(png_bytes("GIF"), "") == "gif" and F.detect(png_bytes("BMP")) == "bmp"
    assert F.detect(png_bytes("WEBP")) == "webp" and F.detect(png_bytes("TIFF")) == "tiff"
    assert F.detect(docx_bytes()) == "docx" and F.detect(pptx_bytes()) == "pptx" and F.detect(xlsx_bytes()) == "xlsx"
    assert F.detect(b"MZ\x90\x00 an exe", "evil.pdf") is None
    assert F.detect(b"", "a.pdf") is None

def add(c, name, data, title=None):
    return c.post("/admin/documents", headers=H(c, "adm"), json={"title": title or name, "filename": name, "content_b64": C.b64e(data)})

def test_rejects_unsupported_and_damaged(c):
    assert add(c, "evil.pdf", b"MZ\x90\x00 not a pdf").status_code == 415
    assert add(c, "broken.png", b"\x89PNG\r\n\x1a\n" + b"junk").status_code == 422
    assert add(c, "empty.pdf", b"").status_code == 415

def release(c, did):
    adm = H(c, "adm")
    kid = c.get("/admin/recipients", headers=adm).json()["recipients"][0]["recipient_key_id"]
    assert c.post("/admin/distributions", headers=adm, json={"document_id": did, "recipient_key_ids": [kid]}).status_code == 200
    r = c.post("/recipient/decrypt-and-release", headers=H(c, "bob"), json={"document_id": did, "keystore_passphrase": "bob-keystore-pass"})
    assert r.status_code == 200, r.text
    return r.json()

def check(c, data, name):
    r = c.post("/submitter/evidence", headers=H(c, "bob"), json={"filename": name, "content_b64": C.b64e(data)})
    assert r.status_code == 200, r.text
    return r.json()

def test_png_roundtrip_and_trace(c):
    d = add(c, "scan.png", png_bytes()); assert d.status_code == 200 and d.json()["kind"] == "png"
    rel = release(c, d.json()["document_id"])
    assert rel["extension"] == "png" and rel["media_type"] == "image/png" and len(rel["page_images"]) == 1
    leaked = C.b64d(rel["protected_b64"])
    assert check(c, leaked, "leak.png")["verdict"] == "verified"
    assert check(c, png_bytes(), "clean.png")["verdict"] != "verified"

def test_jpeg_source_and_gif_converted(c):
    j = add(c, "photo.jpg", png_bytes("JPEG")); assert j.status_code == 200 and j.json()["source_kind"] == "jpeg"
    g = add(c, "anim.gif", png_bytes("GIF")); assert g.status_code == 200 and g.json()["converted"] and g.json()["kind"] == "png"
    assert release(c, g.json()["document_id"])["extension"] == "png"

@needs_office
@pytest.mark.parametrize("name,maker", [("plan.docx", docx_bytes), ("deck.pptx", pptx_bytes), ("sheet.xlsx", xlsx_bytes)])
def test_office_converted_released_as_pdf_and_traced(c, name, maker):
    d = add(c, name, maker()); assert d.status_code == 200, d.text
    assert d.json()["kind"] == "pdf" and d.json()["converted"]
    rel = release(c, d.json()["document_id"])
    assert rel["extension"] == "pdf" and C.b64d(rel["protected_b64"])[:5] == b"%PDF-" and rel["page_images"]
    assert check(c, C.b64d(rel["protected_b64"]), "leak.pdf")["verdict"] == "verified"

def test_doc_list_reports_type(c):
    docs = c.get("/admin/documents", headers=H(c, "adm")).json()["documents"]
    assert {"png", "jpeg", "gif"} <= {d["source_kind"] for d in docs}
    assert c.get("/formats").json()["supported"]

def test_two_shares_in_same_second_both_reach_inbox(c):
    adm = H(c, "adm")
    docs = c.get("/admin/documents", headers=adm).json()["documents"][:2]
    kid = c.get("/admin/recipients", headers=adm).json()["recipients"][0]["recipient_key_id"]
    for d in docs:
        assert c.post("/admin/distributions", headers=adm, json={"document_id": d["document_id"], "recipient_key_ids": [kid]}).status_code == 200
    got = {p["document_id"] for p in c.get("/recipient/packages", headers=H(c, "bob")).json()["packages"]}
    assert {d["document_id"] for d in docs} <= got
